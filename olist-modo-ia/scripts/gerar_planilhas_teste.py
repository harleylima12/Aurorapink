"""Gera as planilhas de teste do Modo Universal em evals/planilhas/ (dados 100% fictícios).

Cada planilha testa um desafio da leitura (seção 7A e 17 da especificação) e o perfil esperado de cada
coluna fica em evals/planilhas/esperado.json. Determinístico (semente fixa): rodar de novo gera os
mesmos bytes.

    python scripts/gerar_planilhas_teste.py

Nomes, CPFs, e-mails e telefones são inventados (CPFs com dígito verificador válido, mas gerados ao acaso).
"""

from __future__ import annotations

import csv
import io
import json
import random
from datetime import date, timedelta
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
PASTA = RAIZ / "evals" / "planilhas"

NOMES = ["Ana", "Bruno", "Carla", "Diego", "Elisa", "Fábio", "Gabriela", "Heitor", "Íris", "João", "Karina", "Lucas",
         "Marina", "Nicolas", "Olívia", "Paulo", "Queila", "Rafael", "Sofia", "Tiago", "Úrsula", "Vitor", "Wanda", "Yasmin"]
SOBRENOMES = ["Silva", "Souza", "Oliveira", "Santos", "Lima", "Pereira", "Costa", "Rodrigues", "Almeida", "Nascimento",
              "Araújo", "Ferreira", "Gomes", "Ribeiro", "Martins", "Carvalho"]
UFS = ["SP", "RJ", "MG", "RS", "PR", "SC", "BA", "PE", "CE", "GO", "DF", "ES"]
CIDADES = {"SP": ["São Paulo", "Campinas", "Santos"], "RJ": ["Rio de Janeiro", "Niterói"], "MG": ["Belo Horizonte", "Uberlândia"],
           "RS": ["Porto Alegre", "Caxias do Sul"], "PR": ["Curitiba", "Londrina"], "SC": ["Florianópolis", "Joinville"],
           "BA": ["Salvador"], "PE": ["Recife"], "CE": ["Fortaleza"], "GO": ["Goiânia"], "DF": ["Brasília"], "ES": ["Vitória"]}
PRODUTOS = {"Eletrônicos": ["Fone Bluetooth", "Carregador USB-C", "Caixa de Som", "Smartwatch"],
            "Casa": ["Luminária", "Jogo de Toalhas", "Panela Antiaderente", "Organizador"],
            "Moda": ["Camiseta Básica", "Tênis Casual", "Mochila", "Boné"],
            "Beleza": ["Protetor Solar", "Hidratante", "Perfume 50 ml"],
            "Esporte": ["Garrafa Térmica", "Tapete de Yoga", "Corda de Pular"]}
PRECOS = {"Eletrônicos": (89, 899), "Casa": (25, 260), "Moda": (29, 349), "Beleza": (19, 219), "Esporte": (15, 180)}


def nome(r: random.Random) -> str:
    return f"{r.choice(NOMES)} {r.choice(SOBRENOMES)}"


def cpf(r: random.Random) -> str:
    n = [r.randint(0, 9) for _ in range(9)]
    for _ in range(2):
        s = sum(v * p for v, p in zip(n, range(len(n) + 1, 1, -1)))
        d = (s * 10) % 11
        n.append(0 if d == 10 else d)
    t = "".join(map(str, n))
    return f"{t[:3]}.{t[3:6]}.{t[6:9]}-{t[9:]}"


def email(n: str, dominio: str) -> str:
    base = n.lower().replace(" ", ".")
    for a, b in zip("áéíóúãõâêôçü", "aeiouaoaeocu"):
        base = base.replace(a, b)
    return f"{base}@{dominio}"


def telefone(r: random.Random) -> str:
    return f"({r.randint(11, 99)}) 9{r.randint(1000, 9999)}-{r.randint(1000, 9999)}"


def br(valor: float, casas: int = 2) -> str:
    """1234.5 -> '1.234,50' (formato brasileiro)."""
    texto = f"{valor:,.{casas}f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def escrever_csv(caminho: Path, linhas: list[list[str]], sep: str = ",", codificacao: str = "utf-8") -> None:
    buffer = io.StringIO()
    csv.writer(buffer, delimiter=sep, lineterminator="\r\n").writerows(linhas)
    caminho.write_bytes(buffer.getvalue().encode(codificacao))


def vendas_br(r: random.Random, ids_clientes: list[str]) -> list[list[str]]:
    """CSV do Excel brasileiro: ';', vírgula decimal, 'R$', dd/mm/aaaa, Latin-1, Sim/Não."""
    cab = ["Data da Venda", "Nº Pedido", "Cliente ID", "Produto", "Categoria", "UF", "Cidade", "Quantidade",
           "Preço Unitário", "Valor Total", "Desconto (%)", "Vendedor", "Pago?"]
    vendedores = [nome(r) for _ in range(6)]
    linhas = [cab]
    inicio = date(2023, 1, 1)
    for i in range(1500):
        cat = r.choice(list(PRODUTOS))
        uf = r.choice(UFS)
        qtd = r.choice([1, 1, 1, 2, 2, 3, 5])
        preco = round(r.uniform(*PRECOS[cat]), 2)
        desc = r.choice([0, 0, 0, 5, 10, 15])
        total = round(qtd * preco * (1 - desc / 100), 2)
        dia = inicio + timedelta(days=r.randint(0, 729))
        linhas.append([dia.strftime("%d/%m/%Y"), f"PED-{10000 + i}", r.choice(ids_clientes), r.choice(PRODUTOS[cat]), cat, uf,
                       r.choice(CIDADES[uf]), str(qtd), f"R$ {br(preco)}", f"R$ {br(total)}", f"{br(desc, 1)}%",
                       r.choice(vendedores), r.choice(["Sim", "Sim", "Sim", "Não"])])
    return linhas


def clientes(r: random.Random) -> tuple[list[list[str]], list[str]]:
    cab = ["Cliente ID", "Nome", "E-mail", "Telefone", "CPF", "UF", "Segmento", "Data de Cadastro"]
    linhas = [cab]
    ids = []
    for i in range(300):
        n = nome(r)
        cid = f"C{2000 + i}"
        ids.append(cid)
        linhas.append([cid, n, email(n, "exemplo.com.br"), telefone(r), cpf(r), r.choice(UFS),
                       r.choice(["Varejo", "Atacado", "Governo", "Pessoa Física"]),
                       (date(2021, 1, 1) + timedelta(days=r.randint(0, 900))).isoformat()])
    return linhas, ids


def financeiro(r: random.Random) -> tuple[list[list[str]], list[list[object]]]:
    """Linha de título, linha em branco, cabeçalho, dados e total (a versão .xlsx tem duas abas)."""
    centros = ["Comercial", "Marketing", "Operações", "TI", "RH"]
    cab = ["Mês", "Centro de Custo", "Receita", "Despesa", "Margem %", "Observação"]
    csv_linhas: list[list[str]] = [["Relatório Financeiro 2023 - Empresa Fictícia Ltda", "", "", "", "", ""],
                                   ["Gerado em 05/01/2024", "", "", "", "", ""], ["", "", "", "", "", ""], cab]
    xlsx_linhas: list[list[object]] = []
    tot_r = tot_d = 0.0
    for mes in range(1, 13):
        for c in centros:
            receita = round(r.uniform(20000, 90000), 2) if c in ("Comercial", "Operações") else round(r.uniform(0, 8000), 2)
            despesa = round(r.uniform(8000, 40000), 2)
            margem = (receita - despesa) / receita if receita else 0
            obs = r.choice(["", "", "Ajuste de fechamento", "Inclui bônus", "Campanha sazonal", "Revisar com a contabilidade"])
            tot_r += receita
            tot_d += despesa
            dia = date(2023, mes, 1)
            csv_linhas.append([dia.strftime("%d/%m/%Y"), c, br(receita), br(despesa), f"{br(margem * 100, 1)}%", obs])
            xlsx_linhas.append([dia, c, receita, despesa, margem, obs])
    csv_linhas.append(["Total", "", br(tot_r), br(tot_d), "", ""])
    return csv_linhas, xlsx_linhas


def rh(r: random.Random) -> list[list[str]]:
    cab = ["Matrícula", "Nome do Colaborador", "CPF", "E-mail corporativo", "Departamento", "Cargo", "Data de Admissão",
           "Salário", "Ativo", "Avaliação", "Horas Extras", "Comentário do gestor"]
    cargos = {"Vendas": ["Vendedor", "Coordenador de Vendas"], "Tecnologia": ["Desenvolvedor", "Analista de Dados", "Tech Lead"],
              "Financeiro": ["Analista Financeiro", "Controller"], "Pessoas": ["Analista de RH", "Recrutador"],
              "Operações": ["Assistente de Logística", "Supervisor"]}
    comentarios = ["Entrega consistente e boa comunicação com o time.", "Precisa melhorar a organização das tarefas.",
                   "Liderou o projeto de migração com ótimo resultado.", "Em adaptação; acompanhar no próximo ciclo.",
                   "Excelente relacionamento com clientes.", "Faltas recorrentes no trimestre; conversar."]
    linhas = [cab]
    for i in range(240):
        n = nome(r)
        dep = r.choice(list(cargos))
        salario = round(r.uniform(2200, 18000), 2)
        linhas.append([str(40000 + i), n, cpf(r), email(n, "empresaficticia.com"), dep, r.choice(cargos[dep]),
                       (date(2015, 1, 1) + timedelta(days=r.randint(0, 3200))).strftime("%d/%m/%Y"), f"R$ {br(salario)}",
                       r.choice(["S", "S", "S", "N"]), br(round(r.uniform(2.5, 5), 1), 1), str(r.randint(0, 40)),
                       r.choice(comentarios) + (f" Ciclo {r.randint(1, 4)}." if r.random() < 0.5 else "")])
    return linhas


def estoque(r: random.Random) -> list[list[str]]:
    """Tab como separador, ponto decimal, datas ISO, true/false."""
    cab = ["SKU", "Descrição", "Categoria", "Fornecedor", "Quantidade em Estoque", "Estoque Mínimo", "Custo Unitário",
           "Última Entrada", "Armazém", "Ativo"]
    fornecedores = ["Alfa Distribuidora", "Beta Importação", "Gama Indústria", "Delta Atacado"]
    linhas = [cab]
    n = 0
    for cat, itens in PRODUTOS.items():
        for item in itens:
            for variante in ["P", "M", "G", "Azul", "Preto", "Branco", "Kit 2", "Kit 3"]:
                n += 1
                linhas.append([f"SKU-{cat[:3].upper()}-{n:04d}", f"{item} {variante}", cat, r.choice(fornecedores),
                               str(r.randint(0, 500)), str(r.choice([5, 10, 20, 50])), f"{r.uniform(*PRECOS[cat]) * 0.45:.2f}",
                               (date(2024, 1, 1) + timedelta(days=r.randint(0, 200))).isoformat(),
                               r.choice(["Guarulhos", "Extrema", "Cajamar"]), r.choice(["true", "true", "false"])])
    return linhas


def baguncada(r: random.Random) -> list[list[str]]:
    """Títulos, linha vazia, nomes de coluna estranhos, linhas vazias no meio, subtotais, formatos misturados."""
    linhas: list[list[str]] = [["CONTROLE DE PEDIDOS - LOJINHA", "", "", "", "", "", "", "", ""],
                               ["(planilha do Zé, não mexer!!)", "", "", "", "", "", "", "", ""],
                               ["", "", "", "", "", "", "", "", ""],
                               ["dt", " cliente ", "produto", "Qtd.", "  VALOR (R$) ", "desc %", "entregue?", "estado", "obs", ""]]
    produtos = ["Bolo de pote", "Brigadeiro (cx 20)", "Torta de limão", "Pão de mel", "Cookie"]
    for mes in (3, 4):
        soma = 0.0
        for _ in range(60):
            dia = date(2024, mes, r.randint(1, 28))
            data = dia.strftime("%d/%m/%Y") if r.random() < 0.7 else dia.isoformat()
            valor = round(r.uniform(12, 180), 2)
            soma += valor
            texto_valor = r.choice([br(valor), f"R$ {br(valor)}", f"R${br(valor)}"])
            linhas.append([data, nome(r), r.choice(produtos), str(r.randint(1, 6)), texto_valor, f"{r.choice([0, 5, 10])}%",
                           r.choice(["Sim", "sim", "Não", "NAO", "S"]), r.choice(["SP", "SP", "RJ", "MG", "sp"]),
                           r.choice(["", "", "", "cliente pediu p/ entregar depois das 18h", "pagou no pix", "retirar na loja"]), ""])
            if r.random() < 0.05:
                linhas.append(["", "", "", "", "", "", "", "", "", ""])
        linhas.append([f"Subtotal {mes:02d}/2024", "", "", "", br(soma), "", "", "", "", ""])
        linhas.append(["", "", "", "", "", "", "", "", "", ""])
    linhas.append(["TOTAL GERAL", "", "", "", "", "", "", "", "", ""])
    return linhas


def desenhada() -> list[list[str]]:
    """Duas tabelas lado a lado, separadas por uma coluna vazia: o app deve AVISAR, não adivinhar."""
    linhas = [["Vendas por loja", "", "", "", "Metas por loja", ""], ["Loja", "Vendas", "", "", "Loja", "Meta"]]
    for loja, v, m in [("Centro", "1.200,00", "1.000,00"), ("Norte", "800,00", "900,00"), ("Sul", "950,00", "1.100,00")]:
        linhas.append([loja, v, "", "", loja, m])
    linhas += [["", "", "", "", "", ""], ["Observações", "", "", "", "", ""], ["Dados de março", "", "", "", "", ""]]
    return linhas


def cegas(r: random.Random) -> dict[str, tuple[list[list[str]], str, str]]:
    """Planilhas escritas DEPOIS das regras do perfil, sem ajustar as regras a elas (resultado da 1ª rodada no D37)."""
    escola = [["Aluno", "Turma", "Nota Final", "Frequência", "Data da Prova", "Aprovado", "E-mail do Responsável", "Parecer"]]
    for i in range(90):
        n = nome(r)
        nota = round(r.uniform(3, 10), 1)
        escola.append([n + f" {r.choice(SOBRENOMES)}", r.choice(["6º A", "6º B", "7º A", "8º C"]), br(nota, 1), f"{r.randint(60, 100)}%",
                       f"{r.randint(1, 28):02d}/{r.choice([3, 6, 9, 11]):02d}/2024", "Sim" if nota >= 6 else "Não", email(n, "familia.com"),
                       r.choice(["Participativo em sala.", "Precisa de reforço em frações.", "Ótimo desempenho no bimestre.", "Faltou às avaliações de recuperação."])])
    chamados = [["Protocolo", "Aberto em", "Prioridade", "Tempo de Resposta (h)", "Atendente", "CNPJ do Cliente", "Resolvido", "Descrição do problema", "Custo do Atendimento"]]
    for i in range(200):
        d = date(2024, 1, 1) + timedelta(days=r.randint(0, 180))
        cnpj = f"{r.randint(10, 99)}.{r.randint(100, 999)}.{r.randint(100, 999)}/0001-{r.randint(10, 99)}"
        chamados.append([f"{2024000 + i}", f"{d.isoformat()} {r.randint(8, 18):02d}:{r.randint(0, 59):02d}", r.choice(["Baixa", "Média", "Alta", "Crítica"]),
                         br(r.uniform(0.2, 48), 1), r.choice(["Equipe N1", "Equipe N2", "Plantão"]), cnpj, r.choice(["true", "false"]),
                         r.choice(["Sistema lento ao emitir nota", "Erro 500 no login", "Dúvida sobre fatura", "Integração parou de sincronizar pedidos"]) + f" (ticket {i})",
                         br(r.uniform(15, 400))])
    us = [["Order Date", "Region", "Revenue", "Units", "Customer Email", "Discount Rate", "Product Line"]]
    for i in range(150):
        d = date(2024, r.randint(1, 12), r.randint(1, 28))
        n = nome(r)
        us.append([d.strftime("%m/%d/%Y"), r.choice(["North", "South", "East", "West"]), f"{r.uniform(100, 5000):,.2f}", str(r.randint(1, 40)),
                   email(n, "mail.com"), f"{r.choice([0, 0.05, 0.1, 0.15]):.2f}", r.choice(["Hardware", "Software", "Services"])])
    return {"cega_escola.csv": (escola, ";", "utf-8"), "cega_chamados.csv": (chamados, ",", "utf-8"), "cega_vendas_en.csv": (us, ",", "utf-8")}


ESPERADO_CEGAS = {
    "cega_escola.csv": {"Aluno": "pessoal", "Turma": "categoria", "Nota Final": "numero", "Frequência": "porcentagem", "Data da Prova": "data",
                        "Aprovado": "booleano", "E-mail do Responsável": "pessoal", "Parecer": "texto"},
    "cega_chamados.csv": {"Protocolo": "id", "Aberto em": "data", "Prioridade": "categoria", "Tempo de Resposta (h)": "numero", "Atendente": "categoria",
                          "CNPJ do Cliente": "pessoal", "Resolvido": "booleano", "Descrição do problema": "texto", "Custo do Atendimento": "dinheiro"},
    "cega_vendas_en.csv": {"Order Date": "data", "Region": "categoria", "Revenue": "dinheiro", "Units": "numero", "Customer Email": "pessoal",
                           "Discount Rate": "porcentagem", "Product Line": "categoria"},
}


def xlsx_financeiro(caminho: Path, linhas: list[list[object]]) -> None:
    from openpyxl import Workbook  # só para gerar o teste; o app lê Excel com SheetJS

    wb = Workbook()
    resumo = wb.active
    resumo.title = "Leia-me"
    resumo.append(["Este arquivo tem os dados na aba 'Dados'."])
    resumo.append(["Fonte: sistema financeiro (fictício)"])
    dados = wb.create_sheet("Dados")
    dados.append(["Relatório Financeiro 2023 - Empresa Fictícia Ltda"])
    dados.merge_cells("A1:F1")
    dados.append([])
    dados.append(["Mês", "Centro de Custo", "Receita", "Despesa", "Margem %", "Observação"])
    for linha in linhas:
        dados.append(linha)
        dados.cell(row=dados.max_row, column=1).number_format = "DD/MM/YYYY"
        for col in (3, 4):
            dados.cell(row=dados.max_row, column=col).number_format = '"R$" #,##0.00'
        dados.cell(row=dados.max_row, column=5).number_format = "0.0%"
    fim = dados.max_row
    dados.append(["Total", None, f"=SUM(C4:C{fim})", f"=SUM(D4:D{fim})", None, None])
    wb.properties.creator = "gerar_planilhas_teste.py"
    # Datas fixas: o .xlsx sai igual a cada execução.
    from datetime import datetime

    wb.properties.created = wb.properties.modified = datetime(2024, 1, 5)
    wb.save(caminho)
    # O .xlsx é um zip: reescreve com data fixa em cada arquivo interno para o resultado ser reprodutível.
    import zipfile

    entradas = []
    with zipfile.ZipFile(caminho) as z:
        entradas = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    with zipfile.ZipFile(caminho, "w", zipfile.ZIP_DEFLATED) as z:
        for nome_arquivo, dados in entradas:
            z.writestr(zipfile.ZipInfo(nome_arquivo, date_time=(2024, 1, 5, 0, 0, 0)), dados, zipfile.ZIP_DEFLATED)


ESPERADO = {
    "vendas_br.csv": {
        "leitura": {"separador": ";", "codificacao": "latin-1", "linha_cabecalho": 0},
        "colunas": {"Data da Venda": "data", "Nº Pedido": "id", "Cliente ID": "id", "Produto": "categoria", "Categoria": "categoria",
                    "UF": "uf", "Cidade": "cidade", "Quantidade": "numero", "Preço Unitário": "dinheiro", "Valor Total": "dinheiro",
                    "Desconto (%)": "porcentagem", "Vendedor": "categoria", "Pago?": "booleano"},
    },
    "clientes.csv": {
        "leitura": {"separador": ",", "codificacao": "utf-8", "linha_cabecalho": 0},
        "colunas": {"Cliente ID": "id", "Nome": "pessoal", "E-mail": "pessoal", "Telefone": "pessoal", "CPF": "pessoal", "UF": "uf",
                    "Segmento": "categoria", "Data de Cadastro": "data"},
    },
    "financeiro_titulo_total.csv": {
        "leitura": {"separador": ",", "codificacao": "utf-8", "linha_cabecalho": 3, "linhas_total_removidas": 1},
        "colunas": {"Mês": "data", "Centro de Custo": "categoria", "Receita": "dinheiro", "Despesa": "dinheiro", "Margem %": "porcentagem",
                    "Observação": "texto"},
    },
    "rh_ficticio.csv": {
        "leitura": {"separador": ",", "codificacao": "utf-8", "linha_cabecalho": 0},
        "colunas": {"Matrícula": "id", "Nome do Colaborador": "pessoal", "CPF": "pessoal", "E-mail corporativo": "pessoal",
                    "Departamento": "categoria", "Cargo": "categoria", "Data de Admissão": "data", "Salário": "dinheiro", "Ativo": "booleano",
                    "Avaliação": "numero", "Horas Extras": "numero", "Comentário do gestor": "texto"},
    },
    "estoque.tsv": {
        "leitura": {"separador": "\t", "codificacao": "utf-8", "linha_cabecalho": 0},
        "colunas": {"SKU": "id", "Descrição": "texto", "Categoria": "categoria", "Fornecedor": "categoria", "Quantidade em Estoque": "numero",
                    "Estoque Mínimo": "numero", "Custo Unitário": "dinheiro", "Última Entrada": "data", "Armazém": "categoria", "Ativo": "booleano"},
    },
    "baguncada.csv": {
        "leitura": {"separador": ",", "codificacao": "utf-8", "linha_cabecalho": 3, "linhas_total_removidas": 3},
        "colunas": {"dt": "data", "cliente": "pessoal", "produto": "categoria", "Qtd.": "numero", "VALOR (R$)": "dinheiro",
                    "desc %": "porcentagem", "entregue?": "booleano", "estado": "uf", "obs": "texto"},
    },
    "desenhada.csv": {"leitura": {"separador": ",", "codificacao": "utf-8", "aviso": "varias_tabelas"}, "colunas": {}},
    "financeiro_titulo_total.xlsx": {
        "leitura": {"aba": "Dados", "linha_cabecalho": 2, "linhas_total_removidas": 1},
        "colunas": {"Mês": "data", "Centro de Custo": "categoria", "Receita": "dinheiro", "Despesa": "dinheiro", "Margem %": "porcentagem",
                    "Observação": "texto"},
        "observacao": "Excel: depende do SheetJS (D39). Enquanto ele não estiver instalado, o teste é pulado.",
    },
}


def main() -> None:
    PASTA.mkdir(parents=True, exist_ok=True)
    r = random.Random(20240105)
    linhas_clientes, ids = clientes(r)
    escrever_csv(PASTA / "clientes.csv", linhas_clientes)
    escrever_csv(PASTA / "vendas_br.csv", vendas_br(r, ids), sep=";", codificacao="latin-1")
    fin_csv, fin_xlsx = financeiro(r)
    escrever_csv(PASTA / "financeiro_titulo_total.csv", fin_csv)
    xlsx_financeiro(PASTA / "financeiro_titulo_total.xlsx", fin_xlsx)
    escrever_csv(PASTA / "rh_ficticio.csv", rh(r))
    linhas_estoque = estoque(r)
    escrever_csv(PASTA / "estoque.tsv", linhas_estoque, sep="\t")
    escrever_csv(PASTA / "baguncada.csv", baguncada(r))
    escrever_csv(PASTA / "desenhada.csv", desenhada())
    # Totais de referência calculados AQUI (Python), para o teste conferir o que o app soma no DuckDB.
    def soma(arquivo: str, coluna: str, sep: str = ",", codificacao: str = "utf-8", pular: int = 0) -> float:
        texto = (PASTA / arquivo).read_bytes().decode(codificacao)
        linhas = list(csv.reader(io.StringIO(texto), delimiter=sep))[pular:]
        i = linhas[0].index(coluna) if coluna in linhas[0] else [c.strip() for c in linhas[0]].index(coluna)
        total = 0.0
        for linha in linhas[1:]:
            if not any(c.strip() for c in linha) or linha[0].strip().lower().startswith(("total", "subtotal")):
                continue
            v = linha[i].replace("R$", "").strip()
            if v:
                total += float(v.replace(".", "").replace(",", ".")) if "," in v or sep == ";" else float(v)
        return round(total, 2)

    ESPERADO["vendas_br.csv"]["totais"] = {"Valor Total": soma("vendas_br.csv", "Valor Total", ";", "latin-1"), "linhas": 1500}
    ESPERADO["rh_ficticio.csv"]["totais"] = {"Salário": soma("rh_ficticio.csv", "Salário"), "linhas": 240}
    ESPERADO["baguncada.csv"]["totais"] = {"VALOR (R$)": soma("baguncada.csv", "VALOR (R$)", pular=3), "linhas": 120}
    ESPERADO["financeiro_titulo_total.csv"]["totais"] = {"Receita": soma("financeiro_titulo_total.csv", "Receita", pular=3), "linhas": 60}
    ESPERADO["estoque.tsv"]["totais"] = {"Custo Unitário": soma("estoque.tsv", "Custo Unitário", "\t"), "linhas": len(linhas_estoque) - 1}
    for arquivo, (linhas, sep, cod) in cegas(random.Random(777)).items():
        escrever_csv(PASTA / arquivo, linhas, sep=sep, codificacao=cod)
        ESPERADO[arquivo] = {"leitura": {"separador": sep, "codificacao": cod, "linha_cabecalho": 0}, "colunas": ESPERADO_CEGAS[arquivo], "cega": True}
    (PASTA / "esperado.json").write_text(json.dumps(ESPERADO, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    total = sum(len(v["colunas"]) for v in ESPERADO.values())
    print(f"{len(ESPERADO)} planilhas em {PASTA.relative_to(RAIZ)} ({total} colunas com perfil esperado)")


if __name__ == "__main__":
    main()
